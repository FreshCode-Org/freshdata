"""``clean_excel`` must preserve zero padding, as ``clean_csv`` already does.

``pandas.read_excel`` infers types exactly as ``read_csv`` does, so a cell that
openpyxl stored as the *text* ``"02134"`` arrived as the integer ``2134`` and
the padding was gone before any cleaning step ran. ``clean_csv`` avoids this
with the ``_csv_io.leading_zero_dtypes`` pre-scan; ``clean_excel`` called
``pd.read_excel`` directly and had no equivalent, so
``preserve_leading_zeros=True`` -- documented as a shared option -- changed
nothing there.

The loss was silent: no warning, no report entry, and a postcode column read as
integers then goes on to be profiled and outlier-checked as a quantity.
"""

from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

import freshdata as fd

ZIPS = ["02134", "10001", "94105", "00501", "07030"]


def _workbook(tmp_path, zips=ZIPS, *, sheet="Sheet1"):
    """Write genuine TEXT cells, so the defect cannot be blamed on the file."""
    path = tmp_path / "zips.xlsx"
    book = openpyxl.Workbook()
    sheet_obj = book.active
    sheet_obj.title = sheet
    sheet_obj.append(["cust", "zip", "qty"])
    for row, zip_code in enumerate(zips, start=1):
        sheet_obj.cell(row=row + 1, column=1, value=f"c{row}")
        cell = sheet_obj.cell(row=row + 1, column=2)
        cell.value = zip_code
        cell.data_type = "s"
        sheet_obj.cell(row=row + 1, column=3, value=row * 10)
    book.save(path)
    return path


def test_the_workbook_really_stores_text(tmp_path):
    """Guard the guard: if the fixture stored numbers, the rest proves nothing."""
    loaded = openpyxl.load_workbook(_workbook(tmp_path))
    assert loaded.active["B2"].value == "02134"
    assert loaded.active["B2"].data_type == "s"


def test_leading_zeros_survive_clean_excel(tmp_path):
    out = fd.clean_excel(_workbook(tmp_path), verbose=False)
    assert out["zip"].tolist() == ZIPS
    assert out["zip"].dtype == object


def test_only_the_padded_column_is_forced_to_text(tmp_path):
    """A genuine quantity column must keep its numeric dtype."""
    out = fd.clean_excel(_workbook(tmp_path), verbose=False)
    assert pd.api.types.is_integer_dtype(out["qty"])


def test_the_csv_and_excel_paths_now_agree(tmp_path):
    """The two entry points are documented as companions; they must match."""
    csv_path = tmp_path / "zips.csv"
    pd.DataFrame({"cust": [f"c{i}" for i in range(1, 6)], "zip": ZIPS}).to_csv(
        csv_path, index=False
    )
    from_csv = fd.clean_csv(csv_path, verbose=False)
    from_excel = fd.clean_excel(_workbook(tmp_path), verbose=False)
    assert from_csv["zip"].tolist() == from_excel["zip"].tolist() == ZIPS


def test_a_csv_to_excel_hand_off_keeps_the_padding(tmp_path):
    """The end-to-end shape that lost data: clean CSV, store as xlsx, clean again."""
    csv_path = tmp_path / "zips.csv"
    xlsx_path = tmp_path / "mid.xlsx"
    pd.DataFrame({"cust": [f"c{i}" for i in range(1, 6)], "zip": ZIPS}).to_csv(
        csv_path, index=False
    )
    cleaned = fd.clean_csv(csv_path, verbose=False)
    pd.DataFrame(cleaned).to_excel(xlsx_path, index=False)
    assert fd.clean_excel(xlsx_path, verbose=False)["zip"].tolist() == ZIPS


def test_preserve_leading_zeros_false_still_opts_out(tmp_path):
    """The option must remain an option, not become unconditional behaviour."""
    out = fd.clean_excel(_workbook(tmp_path), verbose=False, preserve_leading_zeros=False)
    assert pd.api.types.is_integer_dtype(out["zip"])
    assert out["zip"].iloc[0] == 2134


def test_an_explicit_dtype_still_wins(tmp_path):
    """The caller decides types when they say so; the pre-scan must stand down."""
    out = fd.clean_excel(
        _workbook(tmp_path), verbose=False, read_excel_kwargs={"dtype": {"zip": str}}
    )
    assert out["zip"].tolist() == ZIPS


def test_a_column_without_padding_is_untouched(tmp_path):
    """No false positives: ordinary numbers must not be turned into text."""
    out = fd.clean_excel(
        _workbook(tmp_path, zips=["12345", "23456", "34567", "45678", "56789"]),
        verbose=False,
    )
    assert pd.api.types.is_integer_dtype(out["zip"])


def test_a_named_sheet_is_pre_scanned_too(tmp_path):
    """The pre-scan must follow sheet_name, or it samples the wrong sheet."""
    path = _workbook(tmp_path, sheet="Q3")
    out = fd.clean_excel(path, verbose=False, read_excel_kwargs={"sheet_name": "Q3"})
    assert out["zip"].tolist() == ZIPS


def test_selecting_several_sheets_still_raises_the_documented_error(tmp_path):
    """The pre-scan must not mask clean_excel's own multi-sheet rejection."""
    with pytest.raises(TypeError, match="cleans a single sheet"):
        fd.clean_excel(
            _workbook(tmp_path), verbose=False, read_excel_kwargs={"sheet_name": None}
        )
